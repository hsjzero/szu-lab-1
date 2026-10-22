# -*- coding: utf-8 -*-
"""Word 小报构建器：通栏报头 + 双栏正文 + 页面边框 + 插图 + 表格。

内容来自 content/newsletter.json，报头落款支持 {name} {sid} {school} {college} 占位符。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *  # noqa
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import config as cfgmod

BLUE = (0x1F, 0x4E, 0x79)
FULL = Cm(17.4)   # 通栏宽度


def build(cfg, content, out_dir, images_dir="images"):
    doc = new_doc(base_size=10.5)
    sec = doc.sections[0]
    a4(sec, margins=(1.8, 1.8, 1.6, 1.6))
    add_page_border(sec, style="double", sz=8, color="1F4E79", space=8)

    # ---------- 报头 ----------
    mast = content.get("masthead", {})
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t, sz=12, color="1F4E79")
    c = t.rows[0].cells[0]
    c.width = FULL
    set_cell_bg(c, "EAF1F8")
    c.text = ""
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(mast.get("title", ""))
    set_font(r, cn=HEI, en="Arial", size=32, bold=True, color=BLUE)

    if mast.get("pinyin"):
        p = c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(mast["pinyin"])
        set_font(r, en="Arial", size=8, color=(0x60, 0x7D, 0x9B))

    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    add_bottom_line(p, size=6, color="1F4E79")
    org = mast.get("org_line") or "主办：{school}{college}    主编：{name}（学号 {sid}）    {issue}"
    r = p.add_run(cfg.fmt(org.replace("{issue}", mast.get("issue", ""))))
    set_font(r, cn=KAI, size=9, color=(0x40, 0x40, 0x40))

    # ---------- 切双栏（连续分节符，报头保持通栏） ----------
    sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
    a4(sec2, margins=(1.8, 1.8, 1.6, 1.6))
    set_columns(sec2, num=2, space_cm=0.9)
    add_page_border(sec2, style="double", sz=8, color="1F4E79", space=8)

    # ---------- 块渲染 ----------
    def lead(text, size=10.5, cn=SONG, linespace=1.4):
        return para(doc, text, cn=cn, size=size, line=linespace, first_indent=21,
                    space_before=1, space_after=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    def h(title, sub=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        r = p.add_run("▍" + title)
        set_font(r, cn=HEI, size=12.5, bold=True, color=BLUE)
        if sub:
            r = p.add_run("   " + sub)
            set_font(r, cn=KAI, size=9, color=(0x80, 0x80, 0x80))
        return p

    def pic(name, width_cm, caption):
        path = os.path.join(images_dir, name)
        if not os.path.exists(path):
            print("  [warn] 图片缺失，跳过：%s" % path)
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(path, width=Cm(width_cm))
        para(doc, caption, cn=KAI, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER,
             line=1.0, space_before=1, space_after=3, color=(0x70, 0x70, 0x70))

    def drop_cap(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.4
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text[0])
        set_font(r, cn=HEI, size=26, bold=True, color=BLUE)
        r = p.add_run(text[1:])
        set_font(r, cn=SONG, size=10.5)
        return p

    def table(headers, rows, caption):
        t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t, sz=6, color="1F4E79")
        for j, hd in enumerate(headers):
            cell = t.rows[0].cells[j]
            set_cell_bg(cell, "2E75B6")
            set_cell_text(cell, hd, cn=HEI, size=9.5, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, color=(0xFF, 0xFF, 0xFF))
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                set_cell_text(t.rows[i].cells[j], v, cn=SONG, size=9,
                              align=WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.LEFT)
            if i % 2 == 0:
                for j in range(len(headers)):
                    set_cell_bg(t.rows[i].cells[j], "F2F6FA")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(caption)
        set_font(r, cn=KAI, size=8.5, color=(0x70, 0x70, 0x70))

    for blk in content.get("blocks", []):
        kind = blk.get("type")
        if kind == "section":
            h(blk.get("title", ""), blk.get("sub"))
        elif kind == "p":
            text = cfg.fmt(blk.get("text", ""))
            if blk.get("drop_cap") and len(text) > 1:
                drop_cap(text)
            else:
                lead(text)
        elif kind == "figure":
            pic(blk["image"], blk.get("width", 7.8), cfg.fmt(blk.get("caption", "")))
        elif kind == "table":
            table(blk["headers"], blk["rows"], cfg.fmt(blk.get("caption", "")))
        elif kind == "formula":
            para(doc, cfg.fmt(blk.get("text", "")), cn=SONG, en="Cambria Math",
                 size=11, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.2,
                 space_before=2, space_after=2, bold=True, color=BLUE)
        elif kind == "list":
            for item in blk.get("items", []):
                para(doc, "◆ " + cfg.fmt(item), cn=SONG, size=10, line=1.35,
                     left_indent=12, space_before=0, space_after=1)
        elif kind == "editorial":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.4
            add_shading(p, "F5F8FB")
            r = p.add_run(cfg.fmt(blk.get("text", "")))
            set_font(r, cn=KAI, size=10)
            if blk.get("sign"):
                r = p.add_run("\n" + cfg.fmt(blk["sign"]))
                set_font(r, cn=KAI, size=10)
        else:
            print("  [warn] 未知块类型：%r" % kind)

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, cfg.out_name("newsletter"))
    doc.save(path)
    return path


def main(argv):
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--content", default="content")
    ap.add_argument("--out", default="out")
    ap.add_argument("--images", default="images")
    args = ap.parse_args(argv)
    cfg = cfgmod.load(args.config)
    content = cfgmod.load_json(os.path.join(args.content, "newsletter.json"), default={})
    print("OK " + build(cfg, content, args.out, args.images))


if __name__ == "__main__":
    main(sys.argv[1:])
