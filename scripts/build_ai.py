# -*- coding: utf-8 -*-
"""AI 文字处理结果构建器：润色/扩写/简写提示词与输出对比记录。

内容来自 content/ai_tasks.json；每项任务可有多轮"提示词→输出"。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *  # noqa
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import config as cfgmod

MAIN = (0x2E, 0x5C, 0x8A)
RED = (0xC0, 0x50, 0x4D)
GREEN = (0x38, 0x76, 0x3D)


def build(cfg, content, out_dir, images_dir=None):
    doc = new_doc(base_size=11)
    sec = doc.sections[0]
    a4(sec, margins=(2.4, 2.4, 2.2, 2.2))
    ident = cfg.identity
    tool = cfg.fmt(content.get("tool") or cfg.assignment["ai_tool"])

    def title(text, sub=None):
        para(doc, text, cn=HEI, size=17, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1.4,
             space_before=0, space_after=2)
        if sub:
            para(doc, sub, cn=SONG, size=10.5,
                 align=WD_ALIGN_PARAGRAPH.CENTER, line=1.3,
                 space_after=2, color=(0x60, 0x60, 0x60))
        para(doc, "%s%s    %s（%s）    %s" % (ident["school"], ident["college"],
                                              ident["name"], ident["sid"],
                                              cfg.assignment["report_date"]),
             cn=SONG, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER,
             line=1.3, space_after=2)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        add_bottom_line(p, size=12, color="2E5C8A")

    def h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.4
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        set_font(r, cn=HEI, size=14, bold=True, color=MAIN)
        return p

    def h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        set_font(r, cn=HEI, size=12, bold=True)
        return p

    def box(label, text, bg="F2F5F9", label_color=MAIN):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(label)
        set_font(r, cn=HEI, size=10.5, bold=True, color=label_color)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.45
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.right_indent = Cm(0.3)
        add_shading(p, bg)
        r = p.add_run(cfg.fmt(text))
        set_font(r, cn=SONG, size=11)
        return p

    def note(text):
        return para(doc, cfg.fmt(text), cn=KAI, size=10, line=1.4,
                    left_indent=12, space_before=0, space_after=4,
                    color=(0x50, 0x50, 0x50))

    def mk_table(headers, rows, widths=None):
        t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t, sz=6, color="2E5C8A")
        for j, hd in enumerate(headers):
            c = t.rows[0].cells[j]
            set_cell_bg(c, "2E5C8A")
            set_cell_text(c, hd, cn=HEI, size=10, bold=True,
                          align=WD_ALIGN_PARAGRAPH.CENTER, color=(0xFF, 0xFF, 0xFF))
            set_cell_valign(c, "center")
        set_repeat_header(t.rows[0])
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                set_cell_text(t.rows[i].cells[j], v, cn=SONG, size=9.5,
                              align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
                set_cell_valign(t.rows[i].cells[j], "center")
            if i % 2 == 0:
                for j in range(len(headers)):
                    set_cell_bg(t.rows[i].cells[j], "EEF3F9")
        if widths:
            for r_ in t.rows:
                for j, w in enumerate(widths):
                    r_.cells[j].width = Cm(w)
        return t

    # ================= 标题 =================
    title("AI 辅助文字处理实验结果",
          "工具：%s    任务：润色 / 扩写 / 简写" % tool)

    # ================= 实验说明 =================
    h1("一、实验说明")
    para(doc, cfg.fmt(content.get("intro",
                                  "本次实验使用 %s，对原稿分别进行润色、扩写与简写处理，"
                                  "重点观察不同提示词对输出结果的影响。" % tool)),
         cn=SONG, size=11, line=1.5, first_indent=22, space_after=4)

    # ================= 各任务 =================
    cn_num = ["一", "二", "三", "四", "五", "六", "七", "八"]
    for ti, task in enumerate(content.get("tasks", [])):
        h1("%s、任务%s：%s" % (cn_num[(ti + 1) % len(cn_num)], ti + 1, task.get("name", "")))
        if task.get("original"):
            h2("%d.1  原稿" % (ti + 1))
            box("【原文】", task["original"], bg="FBF3F3", label_color=RED)
            if task.get("problems"):
                note("原稿存在的问题：" + task["problems"])
        rounds = task.get("rounds", [])
        for ri, rd in enumerate(rounds, start=2):
            h2("%d.%d  %s" % (ti + 1, ri, rd.get("title", "提示词与输出")))
            box("【提示词 %d】" % ri, rd.get("prompt", ""), bg="E8F1E4", label_color=GREEN)
            box("【输出结果 %d】" % ri, rd.get("output", ""))
            if rd.get("note"):
                note("结果分析：" + rd["note"])
        if task.get("compare"):
            para(doc, cfg.fmt(task["compare"]), cn=SONG, size=11, line=1.5,
                 first_indent=22, space_after=4)

    # ================= 对比表 =================
    st = content.get("summary_table")
    if st:
        h1("五、不同提示词对输出结果的影响")
        mk_table(st["headers"], st["rows"], st.get("widths"))
        if st.get("caption"):
            para(doc, cfg.fmt(st["caption"]), cn=KAI, size=10,
                 align=WD_ALIGN_PARAGRAPH.CENTER, line=1.2,
                 space_before=2, space_after=8)

    # ================= 小结 =================
    if content.get("conclusions"):
        h1("六、实验小结")
        for t in content["conclusions"]:
            para(doc, cfg.fmt(t), cn=SONG, size=11, line=1.5, first_indent=22,
                 left_indent=10, space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    if content.get("disclaimer"):
        para(doc, cfg.fmt(content["disclaimer"]), cn=KAI, size=9.5, line=1.4,
             space_before=8, space_after=2, color=(0x80, 0x80, 0x80))

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, cfg.out_name("ai"))
    doc.save(path)
    return path


def main(argv):
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--content", default="content")
    ap.add_argument("--out", default="out")
    ap.add_argument("--images", default="images")
    args = ap.parse_args(argv)
    cfg = cfgmod.load(args.config)
    content = cfgmod.load_json(os.path.join(args.content, "ai_tasks.json"), default={})
    print("OK " + build(cfg, content, args.out, args.images))


if __name__ == "__main__":
    main(sys.argv[1:])
