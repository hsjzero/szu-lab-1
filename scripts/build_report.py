# -*- coding: utf-8 -*-
"""实验报告构建器：封面 + 目的 + 内容 + 步骤与过程（可嵌截图）+ 体会 + 思考题。

内容来自 content/report_steps.json；截图来自 --shots 目录（缺失时显示占位提示）。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *  # noqa
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import config as cfgmod

CW = Cm(16.0)   # 正文宽度
NAVY = (0x1F, 0x38, 0x64)


def build(cfg, content, out_dir, shots_dir="shots_crop"):
    doc = new_doc(base_size=10.5)
    sec = doc.sections[0]
    a4(sec, margins=(2.5, 2.5, 2.4, 2.2))
    ident = cfg.identity
    asg = cfg.assignment

    def p(text="", size=10.5, cn=SONG, bold=False, indent=0,
          align=WD_ALIGN_PARAGRAPH.LEFT, line=1.5, sb=0, sa=3, color=None):
        return para(doc, text, cn=cn, en="Times New Roman", size=size, bold=bold,
                    align=align, line=line, first_indent=indent,
                    space_before=sb, space_after=sa, color=color)

    def sec_head(text):
        p_ = doc.add_paragraph()
        p_.paragraph_format.space_before = Pt(12)
        p_.paragraph_format.space_after = Pt(4)
        p_.paragraph_format.line_spacing = 1.3
        p_.paragraph_format.keep_with_next = True
        r = p_.add_run(text)
        set_font(r, cn=HEI, size=12, bold=True)
        return p_

    def sub_head(text):
        p_ = doc.add_paragraph()
        p_.paragraph_format.space_before = Pt(8)
        p_.paragraph_format.space_after = Pt(2)
        p_.paragraph_format.line_spacing = 1.3
        p_.paragraph_format.keep_with_next = True
        r = p_.add_run(text)
        set_font(r, cn=HEI, size=10.5, bold=True, color=NAVY)
        return p_

    def shot(name, width_cm, caption):
        path = os.path.join(shots_dir, name)
        p_ = doc.add_paragraph()
        p_.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_.paragraph_format.space_before = Pt(4)
        p_.paragraph_format.space_after = Pt(1)
        p_.paragraph_format.line_spacing = 1.0
        p_.paragraph_format.keep_with_next = True
        if os.path.exists(path):
            p_.add_run().add_picture(path, width=Cm(width_cm))
        else:
            r = p_.add_run("［截图缺失：%s，请补入 %s 目录］" % (name, shots_dir))
            set_font(r, cn=KAI, size=9, color=(0xC0, 0, 0))
        return para(doc, caption, cn=KAI, size=9,
                    align=WD_ALIGN_PARAGRAPH.CENTER, line=1.2,
                    space_before=1, space_after=8, color=(0x40, 0x40, 0x40))

    # ================= 封面 =================
    p("%s 实 验 报 告" % ident["school"], size=22, cn=HEI, bold=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, sb=10, sa=24)
    info = [
        ("课程名称", asg.get("course", ident["course"])),
        ("项目名称", asg.get("title", "")),
        ("学    院", ident["college"]),
        ("专    业", ident["major"]),
        ("指导教师", ident["teacher"]),
        ("报 告 人", "%s        学号：%s" % (ident["name"], ident["sid"])),
        ("实验时间", asg.get("date_range", "")),
        ("提交时间", asg.get("due_date", "")),
    ]
    t = doc.add_table(rows=len(info), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(t)
    for i, (k, v) in enumerate(info):
        c0, c1 = t.rows[i].cells
        c0.width = Cm(4.0)
        c1.width = Cm(9.0)
        set_cell_text(c0, k, cn=HEI, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(c1, "：" + v, cn=SONG, size=12, align=WD_ALIGN_PARAGRAPH.LEFT)
        for c in (c0, c1):
            for pp in c.paragraphs:
                pp.paragraph_format.space_before = Pt(5)
                pp.paragraph_format.space_after = Pt(5)
    p("", sa=0)
    p("", sa=0)
    p("%s 教 务 处 制" % ident["school"], size=12, cn=HEI, bold=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, sa=0)

    # ================= 一、实验目的与要求 =================
    doc.add_page_break()
    sec_head("一、实验目的与要求")
    for i, g in enumerate(content.get("goals", [
            "掌握 Word 的基本操作和排版技术；",
            "掌握图文混排、表格处理、样式和模板技术；",
            "设计并制作个人简历、宣传小报或毕业论文等长文档；",
            "掌握利用 AI 工具提升文字信息处理效果；",
            "学会自己提出问题，并得出解决问题的方法。"]), 1):
        p("%d．%s" % (i, g), indent=21)

    # ================= 二、实验内容与方法 =================
    sec_head("二、实验内容与方法")
    for i, item in enumerate(content.get("tasks_overview", [
            "个人简历：按本人真实情况重新撰写并排版，提交个人简历电子文档。",
            "Word 文档小报：自行设计一篇主题自选的 Word 小报，须包含文字、图像、表格、"
            "页面边框，并含有本人姓名学号等信息，以文章主题命名文件。",
            "毕业论文：根据素材编辑成关于某一主题的一篇论文，主题体现正能量；正文至少三页，"
            "包含图像、表格、页眉、页脚，且从正文开始才有页眉页脚；包括封面、自动目录、"
            "摘要、关键词等。",
            "AI 文档：结合不同的提示词，对输入文本进行写作润色、扩写、简写等操作。",
            "将简历、小报、毕业论文设计和 AI 工具应用中的难点／关键步骤（至少 3 项）截图，"
            "把具体操作方法和操作结果记录在“三、实验步骤与过程”中。"]), 1):
        p("%d．%s" % (i, item), indent=21)
    if content.get("topics"):
        p(cfg.fmt(content["topics"]), indent=21)

    # ================= 三、实验步骤与过程 =================
    sec_head("三、实验步骤与过程")
    if content.get("intro_steps"):
        p(cfg.fmt(content["intro_steps"]), indent=21, sa=6)
    else:
        p("下面记录的是这次实验里真正卡住过我、或者反复调了好几次才做对的地方。"
          "每一步都写了当时的问题、怎么操作、最后是什么结果。", indent=21, sa=6)

    fig_no = 0
    for st in content.get("steps", []):
        sub_head(st.get("title", ""))
        if st.get("problem"):
            p(cfg.fmt("问题：" + st["problem"]), indent=21)
        if st.get("method"):
            p(cfg.fmt("操作方法：" + st["method"]), indent=21)
        for sh in st.get("shots", []):
            fig_no += 1
            caption = cfg.fmt(sh.get("caption", "")) or "图 %d" % fig_no
            shot(sh["image"], sh.get("width", 11.2), caption)
        if st.get("result"):
            p(cfg.fmt("结果：" + st["result"]), indent=21)

    # 小结表
    st_tbl = content.get("summary_table")
    if st_tbl:
        tt = doc.add_table(rows=len(st_tbl["rows"]) + 1, cols=len(st_tbl["headers"]))
        tt.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tt, sz=6, color="000000")
        for j, hd in enumerate(st_tbl["headers"]):
            c = tt.rows[0].cells[j]
            set_cell_bg(c, "D9E2F3")
            set_cell_text(c, hd, cn=HEI, size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_valign(c, "center")
        for i, row in enumerate(st_tbl["rows"], 1):
            for j, v in enumerate(row):
                set_cell_text(tt.rows[i].cells[j], v, cn=SONG, size=9.5,
                              align=WD_ALIGN_PARAGRAPH.CENTER if j < 3 else WD_ALIGN_PARAGRAPH.LEFT)
                set_cell_valign(tt.rows[i].cells[j], "center")
        for r_ in tt.rows:
            r_.cells[0].width = Cm(2.2)
            r_.cells[1].width = Cm(2.4)
            r_.cells[2].width = Cm(2.6)
            if len(tt.columns) > 3:
                r_.cells[3].width = Cm(8.8)
        if st_tbl.get("caption"):
            p(cfg.fmt(st_tbl["caption"]), size=9, cn=KAI,
              align=WD_ALIGN_PARAGRAPH.CENTER, sa=6, color=(0x50, 0x50, 0x50))

    # ================= 四、实验结论或体会 =================
    sec_head("四、实验结论或体会")
    for c in content.get("conclusions", []):
        p(cfg.fmt(c), indent=21, sa=5)

    # ================= 五、思考题 =================
    sec_head("五、思考题")
    for qi, qa in enumerate(content.get("thinking", []), 1):
        p("%d．%s" % (qi, cfg.fmt(qa.get("q", ""))), indent=21, sa=3)
        for a in qa.get("a", []):
            p(cfg.fmt(a), indent=32)

    # ================= 批阅 =================
    p("", sa=0)
    p("", sa=0)
    sec_head("指导教师批阅意见：")
    p("", sa=0)
    for _ in range(4):
        p("", sa=0)
    p("成绩评定：", size=10.5, cn=HEI, bold=True, indent=0, sa=0)
    p("", sa=0)
    p("", sa=0)
    p("指导教师签字：                              年      月      日", indent=0, sa=0,
      align=WD_ALIGN_PARAGRAPH.RIGHT)

    # ================= 备注 =================
    p("", sa=0)
    p("备注：", size=10.5, cn=HEI, bold=True, sa=0)
    p("注：1、报告内的项目或内容设置，可根据实际情况加以调整和补充。",
      size=9, cn=SONG, line=1.3, sa=0)
    p("　　2、教师批改学生实验报告时间应在学生提交实验报告时间后 10 日内。",
      size=9, cn=SONG, line=1.3, sa=0)

    # 页脚页码
    add_page_number_footer(sec, size=9)

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, cfg.out_name("report"))
    doc.save(path)
    return path


def main(argv):
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--content", default="content")
    ap.add_argument("--out", default="out")
    ap.add_argument("--images", default="images")
    ap.add_argument("--shots", default="shots_crop")
    args = ap.parse_args(argv)
    cfg = cfgmod.load(args.config)
    content = cfgmod.load_json(os.path.join(args.content, "report_steps.json"), default={})
    print("OK " + build(cfg, content, args.out, args.shots))


if __name__ == "__main__":
    main(sys.argv[1:])
