# -*- coding: utf-8 -*-
"""交付前自检：页数、节数、域、Heading 样式、占位符残留、兼容性模式。

用法：
    python check_docx.py <docx文件或目录> [更多...]

页数通过本机 Word COM 读取（python-docx 无法判断分页）；
没有 Word 时跳过页数检查，其余照常。
"""
import os
import re
import sys
import zipfile

PLACEHOLDER_PAT = re.compile(
    r"（在 Word 中右键|TODO|XXX|请修改|占位|示例内容|截图缺失")


def iter_docx(paths):
    for p in paths:
        if os.path.isdir(p):
            for f in sorted(os.listdir(p)):
                if f.endswith(".docx") and not f.startswith("~$"):
                    yield os.path.join(p, f)
        elif p.endswith(".docx"):
            yield p


def check_one(path):
    from docx import Document
    doc = Document(path)
    info = []
    info.append("节数=%d" % len(doc.sections))
    heads = {}
    for p in doc.paragraphs:
        s = p.style.name if p.style is not None else ""
        if s.startswith("Heading"):
            heads[s] = heads.get(s, 0) + 1
    info.append("标题样式=%s" % (heads or "无"))
    n_tables = len(doc.tables)
    n_images = len(doc.inline_shapes)
    info.append("表格=%d 图片=%d" % (n_tables, n_images))

    # 域与占位符
    xml = doc.element.xml
    fields = []
    for code in ("TOC", "PAGE", "NUMPAGES", "SECTIONPAGES"):
        if code in xml:
            fields.append(code)
    info.append("域=%s" % (",".join(fields) or "无"))
    hits = PLACEHOLDER_PAT.findall(
        "\n".join(p.text for p in doc.paragraphs) + xml)
    if hits:
        info.append("⚠ 占位/残留=%d 处 %s" % (len(hits), sorted(set(hits))[:5]))

    # 兼容性模式
    with zipfile.ZipFile(path) as z:
        try:
            s = z.read("word/settings.xml").decode("utf-8")
        except KeyError:
            s = ""
    if 'w:name="compatibilityMode"' in s and 'w:val="14"' in s:
        info.append("⚠ compatibilityMode=14（兼容性模式，跑 fix_compat.py）")

    # 页数（需本机 Word）
    try:
        import win32com.client as wc
        app = wc.Dispatch("Word.Application")
        app.Visible = False
        try:
            d = app.Documents.Open(os.path.abspath(path), ReadOnly=True)
            pages = d.ComputeStatistics(2)  # wdStatisticPages
            words = d.ComputeStatistics(0)  # wdStatisticWords
            info.append("页数=%d 字数=%d" % (pages, words))
            d.Close(False)
        finally:
            app.Quit()
    except Exception as e:
        info.append("页数=（跳过：%s）" % type(e).__name__)

    return info


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 1
    for p in iter_docx(argv[1:]):
        print("== %s ==" % os.path.basename(p))
        for line in check_one(p):
            print("   " + line)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
