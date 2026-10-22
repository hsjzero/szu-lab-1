# -*- coding: utf-8 -*-
"""个人简历构建器：config.json + content/resume.json -> 姓名-学号-个人简历.docx"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *  # noqa: F401,F403
from common import (para, set_font, add_shading, add_bottom_line, remove_table_borders,
                    set_cell_text, new_doc, a4, SONG, HEI, KAI)
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import config as cfgmod


def build(cfg, content, out_dir, images_dir=None):
    doc = new_doc(base_size=10.5)
    sec = doc.sections[0]
    a4(sec, margins=(2.2, 2.2, 1.8, 1.8))

    ident = cfg.identity

    def hbar(text, color="2E5C8A", size=11):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        add_shading(p, "DCE6F1")
        r = p.add_run("  " + text)
        set_font(r, cn=HEI, size=size, bold=True, color=(0x2E, 0x5C, 0x8A))
        return p

    def row_kv(left, right, bold_left=False):
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        remove_table_borders(t)
        widths = [Cm(3.6), Cm(12.6)]
        for r_ in t.rows:
            r_.cells[0].width = widths[0]
            r_.cells[1].width = widths[1]
        set_cell_text(t.rows[0].cells[0], cfg.fmt(left), cn=SONG, size=10.5, bold=bold_left)
        set_cell_text(t.rows[0].cells[1], cfg.fmt(right), cn=SONG, size=10.5)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 0.6
        return t

    # ---------- 头部：姓名 + 学校 ----------
    head = doc.add_table(rows=1, cols=2)
    remove_table_borders(head)
    head.rows[0].cells[0].width = Cm(11.5)
    head.rows[0].cells[1].width = Cm(4.8)
    c0 = head.rows[0].cells[0]
    c0.text = ""
    p = c0.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(cfg.fmt(ident["name"]))
    set_font(r, cn=HEI, en="Arial", size=26, bold=True, color=(0x2E, 0x5C, 0x8A))
    p2 = c0.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.0
    r = p2.add_run(ident.get("name_pinyin", ""))
    set_font(r, en="Arial", size=9, color=(0x80, 0x80, 0x80))

    c1 = head.rows[0].cells[1]
    c1.text = ""
    pp = c1.paragraphs[0]
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pp.paragraph_format.space_after = Pt(0)
    r = pp.add_run("%s · %s" % (ident["school"], ident["college"]))
    set_font(r, cn=KAI, size=9, color=(0x80, 0x80, 0x80))

    # 求职意向条
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    add_bottom_line(p, size=12, color="2E5C8A")
    intent = cfg.fmt(content.get("intent") or cfg.resume["intent"])
    r = p.add_run("求职意向：%s      （%s · %s）" % (intent, ident["college"], ident["major"]))
    set_font(r, cn=HEI, size=10.5, bold=True, color=(0x2E, 0x5C, 0x8A))

    # ---------- 基本信息 ----------
    hbar("基本信息")
    info_rows = content.get("info_rows") or [
        ["性    别", ident["gender"], "出生年月", ident["birth"]],
        ["籍    贯", ident["hometown"], "政治面貌", ident["politics"]],
        ["学    历", "本科（在读）", "专    业", ident["major"]],
        ["毕业院校", ident["school"], "所在学院", ident["college"]],
    ]
    info = doc.add_table(rows=len(info_rows), cols=4)
    remove_table_borders(info)
    w = [Cm(2.4), Cm(4.6), Cm(2.4), Cm(7.2)]
    for i, row in enumerate(info_rows):
        for j, v in enumerate(row):
            cell = info.rows[i].cells[j]
            cell.width = w[j]
            is_label = (j % 2 == 0)
            set_cell_text(cell, cfg.fmt(v), cn=SONG, size=10.5,
                          bold=is_label, color=(0x2E, 0x5C, 0x8A) if is_label else None)

    contact = content.get("contact") or {}
    ct = doc.add_table(rows=1, cols=2)
    remove_table_borders(ct)
    ct.rows[0].cells[0].width = Cm(8.4)
    ct.rows[0].cells[1].width = Cm(8.2)
    set_cell_text(ct.rows[0].cells[0], "联系电话：%s" % contact.get("phone", ident["phone"]),
                  cn=SONG, size=10.5)
    set_cell_text(ct.rows[0].cells[1], "电子邮箱：%s" % contact.get("email", ident["email"]),
                  cn=SONG, size=10.5)

    # ---------- 各栏目 ----------
    def kv_section(title, key):
        hbar(title)
        for left, right in content.get(key, []):
            row_kv(left, right)

    kv_section("教育背景", "education")
    kv_section("专业技能", "skills")
    kv_section("校园经历", "campus")
    kv_section("获奖情况", "awards")

    hbar("自我评价")
    para(doc, cfg.fmt(content.get("self_eval", "")),
         cn=SONG, size=10.5, line=1.5, first_indent=21,
         space_before=2, space_after=2)

    note = content.get("footer_note",
                       "说明：本简历由模板生成，提交前请核对全部内容是否为真实信息。")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(cfg.fmt(note))
    set_font(r, cn=KAI, size=8.5, color=(0xA0, 0xA0, 0xA0))

    path = os.path.join(out_dir, cfg.out_name("resume"))
    doc.save(path)
    return path


def main(argv):
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--content", default="content")
    ap.add_argument("--out", default="out")
    ap.add_argument("--images", default="images")
    args = ap.parse_args(argv)
    cfg = cfgmod.load(args.config)
    content = cfgmod.load_json(os.path.join(args.content, "resume.json"), default={})
    os.makedirs(args.out, exist_ok=True)
    path = build(cfg, content, args.out, args.images)
    print("OK " + path)


if __name__ == "__main__":
    main(sys.argv[1:])
