# -*- coding: utf-8 -*-
"""毕业论文构建器：封面 + 诚信声明 + 自动目录 + 摘要 + 正文 + 参考文献。

结构为三节：封面（无页眉页脚）→ 声明+目录 → 摘要+正文（页码从 1 重新编号）。
必须顺序写入内容（add_section 后立刻写），不能先建节再回填。
内容来自 content/thesis.json。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *  # noqa
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import config as cfgmod
from fix_compat import fix_file

HEAD_DEF = {
    'Heading 1': (15, 14, 7),
    'Heading 2': (13, 9, 5),
    'Heading 3': (12, 6, 3),
}


def build(cfg, content, out_dir, images_dir="images"):
    doc = new_doc(base_size=12)
    ident = cfg.identity
    title = cfg.fmt(content.get("title") or cfg.thesis["title"])
    abstract = cfg.fmt(content.get("abstract", ""))
    keywords = cfg.fmt(content.get("keywords") or cfg.thesis["keywords"])

    # 全局样式
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)
    for name, (sz, sb, sa) in HEAD_DEF.items():
        st = doc.styles[name]
        st.font.name = 'Arial'
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.italic = False
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.append(rf)
        rf.set(qn('w:ascii'), 'Arial')
        rf.set(qn('w:hAnsi'), 'Arial')
        rf.set(qn('w:eastAsia'), HEI)
        st.paragraph_format.space_before = Pt(sb)
        st.paragraph_format.space_after = Pt(sa)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.first_line_indent = Pt(0)

    def setup(s):
        a4(s, margins=(2.6, 2.6, 2.4, 2.2))

    def head(text, level=1):
        p = doc.add_paragraph(style='Heading %d' % level)
        r = p.add_run(cfg.fmt(text))
        sz = HEAD_DEF['Heading %d' % level][0]
        set_font(r, cn=HEI, en="Arial", size=sz, bold=True, color=(0, 0, 0))
        return p

    def body_p(text):
        return para(doc, cfg.fmt(text), cn=SONG, size=12, line=1.6, first_indent=24,
                    space_before=0, space_after=3, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    def fig(name, width_cm, caption):
        path = os.path.join(images_dir, name)
        if not os.path.exists(path):
            print("  [warn] 图片缺失，跳过：%s" % path)
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(path, width=Cm(width_cm))
        para(doc, cfg.fmt(caption), cn=KAI, size=10.5,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1.2,
             space_before=1, space_after=8)

    def make_table(headers, rows, widths=None, caption=None, fs=10.5):
        t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t, sz=6, color="000000")
        for j, hd in enumerate(headers):
            c = t.rows[0].cells[j]
            set_cell_bg(c, "D9E2F3")
            set_cell_text(c, hd, cn=HEI, size=fs, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_valign(c, "center")
        set_repeat_header(t.rows[0])
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                set_cell_text(t.rows[i].cells[j], v, cn=SONG, size=fs,
                              align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
                set_cell_valign(t.rows[i].cells[j], "center")
        if widths:
            for r_ in t.rows:
                for j, w in enumerate(widths):
                    r_.cells[j].width = Cm(w)
        if caption:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(cfg.fmt(caption))
            set_font(r, cn=KAI, size=10.5, bold=True)
        return t

    # ============================================================
    # 第一节：封面
    # ============================================================
    sec1 = doc.sections[0]
    setup(sec1)
    para(doc, "", size=12, space_after=0)
    para(doc, "", size=12, space_after=0)
    para(doc, "", size=12, space_after=0)
    para(doc, ident["school"], cn=HEI, size=26, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.6, space_after=6)
    para(doc, "本 科 毕 业 论 文（设计）", cn=HEI, size=22, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.6, space_after=40)

    title_lines = title.split("——") if "——" in title else [title]
    title_lines = [x.strip() for x in title_lines if x.strip()]
    cover = [
        ("题        目", title_lines, 14),
        ("姓        名", [ident["name"]], 15),
        ("专        业", [ident["major"]], 15),
        ("学        院", [ident["college"]], 15),
        ("学        号", [ident["sid"]], 15),
        ("指 导 教 师", [ident["teacher"]], 15),
        ("职        称", [ident.get("teacher_title", "")], 15),
    ]
    ct = doc.add_table(rows=len(cover), cols=2)
    ct.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(ct)
    for i, (k, vs, fs) in enumerate(cover):
        c0, c1 = ct.rows[i].cells
        c0.width = Cm(3.8)
        c1.width = Cm(10.6)
        set_cell_text(c0, k + "：", cn=HEI, size=15, bold=True,
                      align=WD_ALIGN_PARAGRAPH.RIGHT)
        c1.text = ""
        for n, v in enumerate(vs):
            p = c1.paragraphs[0] if n == 0 else c1.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.6
            r = p.add_run(v)
            set_font(r, cn=SONG, size=fs)

    para(doc, "", size=12, space_after=0)
    para(doc, cfg.assignment["report_date"], cn=SONG, size=15,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, space_after=0)

    # ============================================================
    # 第二节：诚信声明 + 目录
    # ============================================================
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    setup(sec2)
    para(doc, "%s本科毕业论文（设计）" % ident["school"], cn=HEI, size=16, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, space_before=0, space_after=6)
    para(doc, "诚 信 声 明", cn=HEI, size=18, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, space_after=20)
    para(doc, "本人郑重声明：所呈交的毕业论文（设计），题目《%s》是本人在指导教师的指导下，"
              "独立进行研究工作所取得的成果。对本文做出重要贡献的个人和集体，"
              "均已在文中以明确方式注明。除此之外，本论文不包含任何其他个人或集体"
              "已经发表或撰写过的作品成果。本人完全意识到本声明的法律结果。" % title,
         cn=SONG, size=14, line=2.0, first_indent=28, space_after=40)
    para(doc, "", size=12, space_after=0)
    para(doc, "", size=12, space_after=0)
    para(doc, "毕业论文（设计）作者签名：", cn=SONG, size=14, line=1.8,
         align=WD_ALIGN_PARAGRAPH.RIGHT, right_indent=60, space_after=14)
    para(doc, "日期：      年    月    日", cn=SONG, size=14, line=1.8,
         align=WD_ALIGN_PARAGRAPH.RIGHT, right_indent=60, space_after=0)

    # 目录另起一页（同节内分页）
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r._element.append(br)
    para(doc, "目    录", cn=HEI, size=18, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, space_before=6, space_after=14)
    add_toc(doc, "1-3", size=12)

    # ============================================================
    # 第三节：摘要 + 正文 + 参考文献（有页眉页脚）
    # ============================================================
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    setup(sec3)
    restart_page_numbering(sec3, 1)

    para(doc, title, cn=HEI, size=16, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.4, space_before=0, space_after=8)
    author_line = cfg.fmt(cfg.thesis.get("author_line") or
                          "%s（%s）  %s" % (ident["college"], ident["major"], ident["name"]))
    para(doc, author_line, cn=SONG, size=12,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.4, space_after=2)
    para(doc, "学号：%s" % ident["sid"], cn=SONG, size=12,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.4, space_after=14)

    p = doc.add_paragraph(style='Heading 1')
    r = p.add_run("摘要")
    set_font(r, cn=HEI, en="Arial", size=15, bold=True, color=(0, 0, 0))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.6
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("【摘要】")
    set_font(r, cn=HEI, size=12, bold=True)
    r = p.add_run(abstract)
    set_font(r, cn=SONG, size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.6
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("【关键词】")
    set_font(r, cn=HEI, size=12, bold=True)
    r = p.add_run(keywords)
    set_font(r, cn=SONG, size=12)

    # ---------------- 正文 ----------------
    for blk in content.get("blocks", []):
        kind = blk.get("type")
        if kind in ("h1", "h2", "h3"):
            head(blk.get("text", ""), int(kind[1]))
        elif kind == "p":
            body_p(blk.get("text", ""))
        elif kind == "figure":
            fig(blk["image"], blk.get("width", 11.0), blk.get("caption", ""))
        elif kind == "table":
            make_table(blk["headers"], blk["rows"], blk.get("widths"),
                       blk.get("caption"), fs=blk.get("font_size", 10.5))
        elif kind == "ref":
            para(doc, blk.get("text", ""), cn=SONG, size=10.5, line=1.5,
                 space_before=0, space_after=2)
        else:
            print("  [warn] 未知块类型：%r" % kind)

    # 参考文献
    refs = content.get("references") or []
    if refs:
        head("参考文献", 1)
        for r_ in refs:
            para(doc, r_, cn=SONG, size=10.5, line=1.5, space_before=0, space_after=2)

    # ============================================================
    # 页眉页脚：仅第三节（正文）有
    # ============================================================
    for s in (sec1, sec2):
        unlink_header_footer(s)
        for p in list(s.header.paragraphs):
            p.text = ""
        for p in list(s.footer.paragraphs):
            p.text = ""

    unlink_header_footer(sec3)
    hp = sec3.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run(title)
    set_font(r, cn=SONG, size=10.5)

    add_section_page_number_footer(sec3, size=10.5, cn=SONG)

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, cfg.out_name("thesis"))
    doc.save(path)
    fix_file(path)
    return path


def main(argv):
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--content", default="content")
    ap.add_argument("--out", default="out")
    ap.add_argument("--images", default="images")
    args = ap.parse_args(argv)
    cfg = cfgmod.load(args.config)
    content = cfgmod.load_json(os.path.join(args.content, "thesis.json"), default={})
    print("OK " + build(cfg, content, args.out, args.images))


if __name__ == "__main__":
    main(sys.argv[1:])
